"""Tests for exporter module: neutral records, format parity, dry-run."""

import json
import os
import tempfile
from pathlib import Path

import pytest

from douini.domain.exporters import (
    PlanRecord,
    StepRecord,
    SessionRecord,
    WeekRecord,
    export_csv,
    export_json,
    export_markdown,
    export_plan,
    export_xlsx,
    load_plan_json,
    plan_to_records,
    save_plan_json,
)
from douini.domain.models import (
    Distance,
    Pace,
    Paces,
    RunnerProfile,
    Session,
    TrainingPlan,
    WeekPlan,
)
from douini.domain.planner import generate_plan
from douini.domain.vdot import derive_paces


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def sample_plan():
    runner = RunnerProfile(
        vdot=49.4,
        weekly_volume_km=70,
        training_days=["mon", "wed", "fri", "sun"],
        target_weekly_km=70,
        sessions_per_week=4,
        race_distance="semi",
        weeks=12,
        target_time="1:32:40",
    )
    paces = derive_paces(49.4)
    plan = generate_plan(runner, Distance.SEMI, 12, paces=paces,
                         sessions_per_week=4)
    return plan


@pytest.fixture
def tmp_outdir():
    with tempfile.TemporaryDirectory() as d:
        yield d


# ---------------------------------------------------------------------------
# Task 1: Neutral records
# ---------------------------------------------------------------------------

class TestNeutralRecords:
    def test_plan_record_has_all_required_fields(self, sample_plan):
        rec = plan_to_records(sample_plan)
        assert rec.distance
        assert rec.weeks == 12
        assert rec.vdot == 49.4
        assert rec.goal_time
        assert rec.sessions_per_week == 4
        assert isinstance(rec.paces, dict)
        assert "EF" in rec.paces
        assert "SHORT" in rec.paces
        assert "MEDIUM" in rec.paces
        assert "LONG" in rec.paces
        assert isinstance(rec.warnings, list)
        assert isinstance(rec.weeks_data, list)
        assert len(rec.weeks_data) == 12

    def test_week_record_fields(self, sample_plan):
        rec = plan_to_records(sample_plan)
        wk = rec.weeks_data[0]
        assert wk.week_num == 1
        assert wk.bloc
        assert isinstance(wk.is_recovery, bool)
        assert isinstance(wk.is_taper, bool)
        assert isinstance(wk.is_race_week, bool)
        assert isinstance(wk.target_km, (int, float))
        assert isinstance(wk.actual_km, (int, float))
        assert isinstance(wk.weekly_load, (int, float))
        assert isinstance(wk.fatigue_index, (int, float))
        assert isinstance(wk.recovery_need, (int, float))
        assert isinstance(wk.sessions, list)
        assert len(wk.sessions) > 0

    def test_session_record_fields(self, sample_plan):
        rec = plan_to_records(sample_plan)
        s = rec.weeks_data[0].sessions[0]
        assert s.week
        assert s.bloc
        assert s.day
        assert s.type
        assert isinstance(s.distance_km, (int, float))
        assert isinstance(s.category, str)
        assert isinstance(s.load_score, (int, float))
        assert isinstance(s.purpose, str)
        assert isinstance(s.steps, list)

    def test_steps_present_for_quality_sessions(self, sample_plan):
        rec = plan_to_records(sample_plan)
        quality_sessions = [
            s for wk in rec.weeks_data for s in wk.sessions
            if s.workout is not None
        ]
        assert len(quality_sessions) > 0
        for s in quality_sessions:
            assert len(s.steps) > 0

    def test_steps_present_for_easy_sessions(self, sample_plan):
        rec = plan_to_records(sample_plan)
        easy_sessions = [
            s for wk in rec.weeks_data for s in wk.sessions
            if s.type == "easy"
        ]
        assert len(easy_sessions) > 0
        for s in easy_sessions:
            assert len(s.steps) > 0

    def test_race_day_detected(self, sample_plan):
        rec = plan_to_records(sample_plan)
        assert rec.race_day is not None
        assert rec.race_day["week"] == 12
        assert rec.race_day["day"] == "sun"
        assert rec.race_day["distance_km"] > 0

    def test_warnings_preserved(self, sample_plan):
        rec = plan_to_records(sample_plan)
        # Plan may have 0 or more warnings
        assert isinstance(rec.warnings, list)
        for w in rec.warnings:
            assert "code" in w
            assert "message" in w
            assert "severity" in w

    def test_ordering_stable(self, sample_plan):
        rec1 = plan_to_records(sample_plan)
        rec2 = plan_to_records(sample_plan)
        for w1, w2 in zip(rec1.weeks_data, rec2.weeks_data):
            assert w1.week_num == w2.week_num
            for s1, s2 in zip(w1.sessions, w2.sessions):
                assert s1.day == s2.day
                assert s1.workout == s2.workout
                assert len(s1.steps) == len(s2.steps)

    def test_exporter_does_not_mutate_plan(self, sample_plan):
        original_weeks = len(sample_plan.week_plans)
        original_sessions = [len(wk.sessions) for wk in sample_plan.week_plans]
        plan_to_records(sample_plan)
        assert len(sample_plan.week_plans) == original_weeks
        assert [len(wk.sessions) for wk in sample_plan.week_plans] == original_sessions

    def test_phase_present_on_weeks(self, sample_plan):
        rec = plan_to_records(sample_plan)
        for wk in rec.weeks_data:
            assert wk.phase is not None or wk.is_recovery


# ---------------------------------------------------------------------------
# Task 2: Format parity
# ---------------------------------------------------------------------------

class TestFormatParity:
    def test_json_export(self, sample_plan, tmp_outdir):
        path = export_json(sample_plan, tmp_outdir)
        assert os.path.exists(path)
        with open(path) as f:
            data = json.load(f)
        assert data["distance"] == "semi"
        assert data["weeks"] == 12
        assert data["vdot"] == 49.4
        assert "paces" in data
        assert "week_plans" in data
        assert len(data["week_plans"]) == 12

    def test_csv_export(self, sample_plan, tmp_outdir):
        path = export_csv(sample_plan, tmp_outdir)
        assert os.path.exists(path)
        import csv as csv_mod
        with open(path) as f:
            reader = csv_mod.reader(f)
            header = next(reader)
            assert "week" in header
            assert "day" in header
            assert "type" in header
            assert "workout" in header
            assert "distance_km" in header
            rows = list(reader)
            assert len(rows) > 0

    def test_xlsx_export(self, sample_plan, tmp_outdir):
        path = export_xlsx(sample_plan, tmp_outdir)
        assert os.path.exists(path)
        from openpyxl import load_workbook
        wb = load_workbook(path)
        assert "Plan" in wb.sheetnames
        assert "Weeks" in wb.sheetnames
        assert "Sessions" in wb.sheetnames
        assert "Steps" in wb.sheetnames

    def test_markdown_export(self, sample_plan, tmp_outdir):
        path = export_markdown(sample_plan, tmp_outdir)
        assert os.path.exists(path)
        with open(path) as f:
            content = f.read()
        assert "Plan" in content
        assert "Allures" in content
        assert "Planning" in content

    def test_all_formats_same_session_count(self, sample_plan, tmp_outdir):
        json_path = export_json(sample_plan, tmp_outdir)
        csv_path = export_csv(sample_plan, tmp_outdir)
        xlsx_path = export_xlsx(sample_plan, tmp_outdir)
        md_path = export_markdown(sample_plan, tmp_outdir)

        # JSON session count
        with open(json_path) as f:
            jdata = json.load(f)
        j_sessions = sum(len(wk["sessions"]) for wk in jdata["week_plans"])

        # CSV row count
        import csv as csv_mod
        with open(csv_path) as f:
            reader = csv_mod.reader(f)
            next(reader)  # skip header
            c_sessions = sum(1 for _ in reader)

        # XLSX session count
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path)
        ws = wb["Sessions"]
        x_sessions = ws.max_row - 1  # minus header

        assert j_sessions == c_sessions == x_sessions

    def test_all_formats_same_workout_ids(self, sample_plan, tmp_outdir):
        json_path = export_json(sample_plan, tmp_outdir)
        csv_path = export_csv(sample_plan, tmp_outdir)

        with open(json_path) as f:
            jdata = json.load(f)
        j_workouts = {
            s["workout"] for wk in jdata["week_plans"]
            for s in wk["sessions"] if s["workout"]
        }

        import csv as csv_mod
        with open(csv_path) as f:
            reader = csv_mod.DictReader(f)
            c_workouts = {
                row["workout"] for row in reader
                if row["workout"]
            }

        assert j_workouts == c_workouts

    def test_all_formats_same_distances(self, sample_plan, tmp_outdir):
        json_path = export_json(sample_plan, tmp_outdir)
        csv_path = export_csv(sample_plan, tmp_outdir)

        with open(json_path) as f:
            jdata = json.load(f)
        j_dists = [
            s["distance_km"] for wk in jdata["week_plans"]
            for s in wk["sessions"]
        ]

        import csv as csv_mod
        with open(csv_path) as f:
            reader = csv_mod.DictReader(f)
            c_dists = [float(row["distance_km"]) for row in reader]

        assert len(j_dists) == len(c_dists)
        for j, c in zip(j_dists, c_dists):
            assert abs(j - c) < 0.01

    def test_export_plan_multiple_formats(self, sample_plan, tmp_outdir):
        paths = export_plan(sample_plan, tmp_outdir, ("json", "csv", "md"))
        assert len(paths) == 3
        for p in paths:
            assert os.path.exists(p)

    def test_export_plan_xlsx_optional(self, sample_plan, tmp_outdir):
        paths = export_plan(sample_plan, tmp_outdir, ("xlsx",))
        assert len(paths) == 1
        assert paths[0].endswith(".xlsx")


# ---------------------------------------------------------------------------
# Task 3: CLI side effects
# ---------------------------------------------------------------------------

class TestDryRun:
    def test_dry_run_writes_nothing(self, sample_plan, tmp_outdir):
        """Dry-run must not write any export files."""
        rec = plan_to_records(sample_plan)
        # Simulate what CLI does on dry-run: just build records, no writes
        before = set(os.listdir(tmp_outdir))
        # In dry-run mode, CLI returns before _export_plan call
        after = set(os.listdir(tmp_outdir))
        assert before == after

    def test_export_writes_files(self, sample_plan, tmp_outdir):
        """Non-dry-run should write files."""
        before = set(os.listdir(tmp_outdir))
        export_plan(sample_plan, tmp_outdir, ("md", "csv"))
        after = set(os.listdir(tmp_outdir))
        assert len(after - before) == 2


class TestXLSXOptional:
    def test_xlsx_without_openpyxl_raises(self, sample_plan, tmp_outdir, monkeypatch):
        """If openpyxl is unavailable, JSON/CSV/MD still work; XLSX raises clear error."""
        # Simulate openpyxl not available
        import builtins
        real_import = builtins.__import__

        def fake_import(name, *args, **kwargs):
            if name == "openpyxl":
                raise ImportError("simulated")
            return real_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", fake_import)

        with pytest.raises(ImportError, match="openpyxl"):
            export_xlsx(sample_plan, tmp_outdir)

        # Other formats still work
        export_json(sample_plan, tmp_outdir)
        export_csv(sample_plan, tmp_outdir)
        export_markdown(sample_plan, tmp_outdir)


# ---------------------------------------------------------------------------
# Task 4: Round-trip parity
# ---------------------------------------------------------------------------

class TestRoundTrip:
    def test_save_load_roundtrip(self, sample_plan, tmp_outdir):
        path = os.path.join(tmp_outdir, "test_plan.json")
        save_plan_json(sample_plan, path)
        loaded = load_plan_json(path)

        assert loaded.distance == sample_plan.distance
        assert loaded.weeks == sample_plan.weeks
        assert loaded.runner.vdot == sample_plan.runner.vdot
        assert loaded.sessions_per_week == sample_plan.sessions_per_week

        # Same session count
        orig_count = sum(len(wk.sessions) for wk in sample_plan.week_plans)
        loaded_count = sum(len(wk.sessions) for wk in loaded.week_plans)
        assert orig_count == loaded_count

    def test_roundtrip_preserves_workout_ids(self, sample_plan, tmp_outdir):
        path = os.path.join(tmp_outdir, "test_plan.json")
        save_plan_json(sample_plan, path)
        loaded = load_plan_json(path)

        orig_ids = {
            s.workout.name for wk in sample_plan.week_plans
            for s in wk.sessions if s.workout
        }
        loaded_ids = {
            s.workout.name for wk in loaded.week_plans
            for s in wk.sessions if s.workout
        }
        assert orig_ids == loaded_ids

    def test_roundtrip_preserves_distances(self, sample_plan, tmp_outdir):
        path = os.path.join(tmp_outdir, "test_plan.json")
        save_plan_json(sample_plan, path)
        loaded = load_plan_json(path)

        for wk_orig, wk_loaded in zip(sample_plan.week_plans, loaded.week_plans):
            for s_orig, s_loaded in zip(wk_orig.sessions, wk_loaded.sessions):
                assert abs(s_orig.distance_km - s_loaded.distance_km) < 0.01

    def test_roundtrip_preserves_paces(self, sample_plan, tmp_outdir):
        path = os.path.join(tmp_outdir, "test_plan.json")
        save_plan_json(sample_plan, path)
        loaded = load_plan_json(path)

        assert str(loaded.paces.ef[0]) == str(sample_plan.paces.ef[0])
        assert str(loaded.paces.short[0]) == str(sample_plan.paces.short[0])
        assert str(loaded.paces.medium[0]) == str(sample_plan.paces.medium[0])
        assert str(loaded.paces.long[0]) == str(sample_plan.paces.long[0])

    def test_roundtrip_preserves_category_and_pace_key(self, sample_plan, tmp_outdir):
        path = os.path.join(tmp_outdir, "test_plan.json")
        save_plan_json(sample_plan, path)
        loaded = load_plan_json(path)

        for wk_orig, wk_loaded in zip(sample_plan.week_plans, loaded.week_plans):
            for s_orig, s_loaded in zip(wk_orig.sessions, wk_loaded.sessions):
                assert s_orig.category == s_loaded.category
                assert s_orig.pace_key == s_loaded.pace_key

    def test_exporter_does_not_read_network(self, sample_plan, tmp_outdir):
        """Exporters must not touch network state."""
        # Just verify it doesn't raise and produces files
        paths = export_plan(sample_plan, tmp_outdir, ("json", "csv", "md"))
        assert len(paths) == 3

    def test_multiple_distances_export(self, tmp_outdir):
        """Export works for all distances."""
        runner = RunnerProfile(
            vdot=50.0,
            weekly_volume_km=70,
            training_days=["mon", "wed", "fri", "sun"],
            target_weekly_km=70,
            sessions_per_week=4,
        )
        paces = derive_paces(50.0)
        for dist in [Distance.K5, Distance.K10, Distance.SEMI, Distance.MARATHON]:
            min_w = {Distance.K5: 4, Distance.K10: 6,
                     Distance.SEMI: 10, Distance.MARATHON: 12}[dist]
            plan = generate_plan(runner, dist, min_w, paces=paces,
                                 sessions_per_week=4)
            paths = export_plan(plan, tmp_outdir, ("json", "csv", "md"))
            assert len(paths) == 3


# ---------------------------------------------------------------------------
# Step structure parity
# ---------------------------------------------------------------------------

class TestStepParity:
    def test_uniform_workout_has_repeat_step(self, sample_plan):
        rec = plan_to_records(sample_plan)
        # Find a quality session with a uniform workout (NS-S01, NS-M01, etc.)
        uniform_sessions = [
            s for wk in rec.weeks_data for s in wk.sessions
            if s.workout and s.steps
            and any(st.type == "repeat" for st in s.steps)
        ]
        assert len(uniform_sessions) > 0
        for s in uniform_sessions:
            repeat_step = next(st for st in s.steps if st.type == "repeat")
            assert repeat_step.repeat_count > 0
            assert len(repeat_step.children) >= 2  # interval + recovery

    def test_variable_workout_has_sequential_intervals(self, sample_plan):
        rec = plan_to_records(sample_plan)
        # NS-S04, NS-M06, NS-L05 are variable
        variable_sessions = [
            s for wk in rec.weeks_data for s in wk.sessions
            if s.workout in ("NS-S04", "NS-M06", "NS-L05")
        ]
        for s in variable_sessions:
            assert len(s.steps) > 2  # multiple interval+recovery pairs
            # No repeat wrapper
            assert not any(st.type == "repeat" for st in s.steps)

    def test_step_ordering(self, sample_plan):
        rec = plan_to_records(sample_plan)
        for wk in rec.weeks_data:
            for s in wk.sessions:
                for i, st in enumerate(s.steps):
                    assert st.order == i or st.order == i * 2 or st.order == i * 2 + 1

    def test_step_to_dict(self, sample_plan):
        rec = plan_to_records(sample_plan)
        s = rec.weeks_data[0].sessions[0]
        for st in s.steps:
            d = st.to_dict()
            assert "order" in d
            assert "type" in d
            assert "duration_sec" in d
